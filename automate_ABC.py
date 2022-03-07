import openpyxl

path = f'./Curva ABC.xlsx'
book = openpyxl.load_workbook(path) 

sheet = book['Sheet1']

filial = 101

for row in sheet:

    if row[1].value == None:
        sheet.delete_rows(row[0].row, 1)
        sheet.cell(row[0].row, 1).value = filial
    if row[1].value == 'Total Filial:':
        sheet.delete_rows(row[0].row, 3)
        filial += 1
        sheet.cell(row[0].row, 1).value = filial
    row[0].value = filial
        

path = f'./Curva ABC_add_test.xlsx'
book.save(path)
